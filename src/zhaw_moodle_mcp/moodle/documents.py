"""Plain-text extraction from downloaded course files (PDF, PowerPoint, Word, HTML, text),
so that clients without file access (Claude Desktop chat) can read them."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup

from ..errors import ErrorCode, MoodleError

log = logging.getLogger(__name__)

PAGED = {".pdf": "pdf", ".pptx": "powerpoint"}
WHOLE = {".docx": "word", ".html": "html", ".htm": "html", ".txt": "text", ".md": "text", ".csv": "text",
         ".java": "text", ".py": "text", ".json": "text", ".xml": "text", ".sql": "text"}


@dataclass
class DocumentText:
    file_type: str
    text: str
    total_pages: int | None  # pages or slides; None for documents without pages
    first_page: int | None = None
    last_page: int | None = None
    truncated: bool = False
    note: str | None = None


def parse_page_range(spec: str | None, total: int) -> tuple[int, int]:
    """'3', '2-5', '4-' (to the end) -> 1-based inclusive range within 1..total."""
    if not spec or not spec.strip():
        return 1, total
    m = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d*)\s*)?", spec)
    if not m:
        raise MoodleError(ErrorCode.MOODLE_ERROR, f"Invalid pages {spec!r}; use e.g. '5', '1-10' or '11-'.")
    first = int(m.group(1))
    last = first if m.group(2) is None else (int(m.group(2)) if m.group(2) else total)
    if first < 1 or first > total or last < first:
        raise MoodleError(ErrorCode.MOODLE_ERROR, f"Pages {spec!r} are outside 1-{total}.")
    return first, min(last, total)


def _pdf_pages(path: Path) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - broken pages must not stop the rest
            log.info("PDF page could not be read: %s", type(exc).__name__)
            pages.append("")
    return pages


def _pptx_slides(path: Path) -> list[str]:
    from pptx import Presentation

    slides = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notizen: " + slide.notes_slide.notes_text_frame.text.strip())
        slides.append("\n".join(parts))
    return slides


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") and style[-1:].isdigit():
            text = "#" * min(int(style[-1]) + 1, 6) + " " + text
        elif "list" in style:
            text = "- " + text
        parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
    return "\n\n".join(parts)


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+\n", "\n", text.replace("\r", ""))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_text(path: Path, pages: str | None, max_chars: int, base_url: str = "") -> DocumentText:
    suffix = path.suffix.lower()
    if suffix in PAGED:
        file_type = PAGED[suffix]
        label = "Seite" if file_type == "pdf" else "Folie"
        try:
            all_pages = _pdf_pages(path) if file_type == "pdf" else _pptx_slides(path)
        except Exception as exc:  # noqa: BLE001 - corrupt or protected files
            raise MoodleError(ErrorCode.DOWNLOAD_FAILED, f"{path.name} could not be read ({type(exc).__name__}).") \
                from exc
        total = len(all_pages)
        if total == 0:
            return DocumentText(file_type, "", 0, note="The file has no pages.")
        first, last = parse_page_range(pages, total)
        chunks, size, returned_last, truncated = [], 0, first - 1, False
        for number in range(first, last + 1):
            chunk = f"--- {label} {number} ---\n{_clean(all_pages[number - 1])}"
            if chunks and size + len(chunk) > max_chars:
                truncated = True
                break
            chunks.append(chunk[:max_chars])
            size += len(chunk)
            returned_last = number
        text = "\n\n".join(chunks)
        result = DocumentText(file_type, text, total, first, returned_last, truncated)
        if file_type == "pdf" and not any(_clean(p) for p in all_pages[first - 1:returned_last]):
            result.note = "No text found: the PDF is probably scanned images."
        return result

    if suffix in WHOLE:
        file_type = WHOLE[suffix]
        try:
            if file_type == "word":
                text = _docx_text(path)
            else:
                raw = path.read_text(encoding="utf-8", errors="replace")
                if file_type == "html":
                    from .content import to_markdown

                    soup = BeautifulSoup(raw, "lxml")
                    text = to_markdown(soup.body or soup, base_url).text
                else:
                    text = raw
        except Exception as exc:  # noqa: BLE001
            raise MoodleError(ErrorCode.DOWNLOAD_FAILED, f"{path.name} could not be read ({type(exc).__name__}).") \
                from exc
        text = _clean(text)
        return DocumentText(file_type, text[:max_chars], None, truncated=len(text) > max_chars)

    raise MoodleError(
        ErrorCode.RESOURCE_NOT_DOWNLOADABLE,
        f"Text cannot be extracted from {suffix or 'this'} files ({path.name}); "
        "supported: PDF, PowerPoint (.pptx), Word (.docx), HTML and text files.",
    )
