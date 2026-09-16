import asyncio
import io
from pathlib import Path

import pytest
import respx
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from zhaw_moodle_mcp.errors import ErrorCode, MoodleError
from zhaw_moodle_mcp.moodle.documents import extract_text, parse_page_range
from zhaw_moodle_mcp.prompts import STUDY_GUIDE
from zhaw_moodle_mcp.server import create_server
from zhaw_moodle_mcp.service import MoodleService

from .fake_moodle import FakeFile, FakeMoodle


def make_pdf(pages: list[str]) -> bytes:
    """Minimal PDF with one line of Helvetica text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for text in pages:
        page_id, content_id = len(objects) + 1, len(objects) + 2
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                       f"/Resources << /Font << /F1 {len(pages) * 2 + 3} 0 R >> >> >>")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_pptx(slides: list[tuple[str, str]]) -> bytes:
    prs = Presentation()
    for title, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text_frame.text = f"Inhalt {title}"
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Aufgabe 1", level=1)
    doc.add_paragraph("Erklären Sie den Begriff Vertrag.")
    doc.add_paragraph("Offerte", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Punkte"
    table.rows[0].cells[1].text = "10"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(("spec", "expected"), [
    (None, (1, 10)), ("3", (3, 3)), ("2-5", (2, 5)), ("8-", (8, 10)), (" 4 - 20 ", (4, 10)),
])
def test_page_range(spec, expected):
    assert parse_page_range(spec, 10) == expected


@pytest.mark.parametrize("spec", ["0", "11", "5-2", "a", "1,2"])
def test_page_range_rejects(spec):
    with pytest.raises(MoodleError):
        parse_page_range(spec, 10)


def test_pdf_pages_and_truncation(tmp_path):
    pdf = tmp_path / "v.pdf"
    pdf.write_bytes(make_pdf(["Einleitung OOP", "Klassen und Objekte", "Vererbung"]))
    doc = extract_text(pdf, None, 10_000)
    assert (doc.file_type, doc.total_pages, doc.first_page, doc.last_page, doc.truncated) == ("pdf", 3, 1, 3, False)
    assert "--- Seite 2 ---\nKlassen und Objekte" in doc.text

    part = extract_text(pdf, "2-", 45)
    assert (part.first_page, part.last_page, part.truncated) == (2, 2, True)
    assert "Vererbung" not in part.text


def test_pdf_without_text_gets_a_note(tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(make_pdf([""]))
    assert "scanned" in extract_text(pdf, None, 1000).note


def test_pptx_slides_with_notes(tmp_path):
    f = tmp_path / "s.pptx"
    f.write_bytes(make_pptx([("UML", "Wichtig für MEP"), ("Java", "")]))
    doc = extract_text(f, "1", 10_000)
    assert doc.file_type == "powerpoint" and doc.total_pages == 2
    assert doc.text == "--- Folie 1 ---\nUML\nInhalt UML\nNotizen: Wichtig für MEP"


def test_docx_html_text(tmp_path):
    (tmp_path / "a.docx").write_bytes(make_docx())
    assert extract_text(tmp_path / "a.docx", None, 10_000).text == (
        "## Aufgabe 1\n\nErklären Sie den Begriff Vertrag.\n\n- Offerte\n\n| Punkte | 10 |")
    (tmp_path / "p.html").write_text("<html><body><h1>Quiz</h1><p>Frage <b>1</b></p></body></html>",
                                     encoding="utf-8")
    assert extract_text(tmp_path / "p.html", None, 10_000).text == "## Quiz\n\nFrage **1**"
    (tmp_path / "t.txt").write_text("x" * 50, encoding="utf-8")
    doc = extract_text(tmp_path / "t.txt", None, 20)
    assert doc.text == "x" * 20 and doc.truncated and doc.total_pages is None


def test_unsupported_and_broken_files(tmp_path):
    (tmp_path / "a.zip").write_bytes(b"PK")
    with pytest.raises(MoodleError) as exc:
        extract_text(tmp_path / "a.zip", None, 100)
    assert exc.value.code == ErrorCode.RESOURCE_NOT_DOWNLOADABLE
    (tmp_path / "bad.pptx").write_bytes(b"not a pptx")
    with pytest.raises(MoodleError) as exc:
        extract_text(tmp_path / "bad.pptx", None, 100)
    assert exc.value.code == ErrorCode.DOWNLOAD_FAILED


@pytest.fixture
def moodle():
    fake = FakeMoodle(
        resources={
            11: ("Vorlesung 1", "Vorlesung-01.pdf", FakeFile(make_pdf(["Titel", "Inhalt"]))),
            12: ("Folien", "Folien.pptx", FakeFile(make_pptx([("Agenda", "")]))),
        },
        folder={"a.txt": FakeFile(b"A"), "b.txt": FakeFile(b"B")},
    )
    with respx.mock(assert_all_called=False) as router:
        fake.mount(router)
        yield fake


@pytest.fixture
async def service(config, store):
    svc = MoodleService(config)
    yield svc
    await svc.aclose()


async def test_read_file_downloads_once(service, moodle, config):
    first = await service.read_file("11", None, 30_000)
    assert first.file_type == "pdf" and first.total_pages == 2 and first.pages == "1-2"
    assert "--- Seite 2 ---\nInhalt" in first.text
    assert Path(first.path).is_relative_to(config.moodle.download_directory)
    downloads = moodle.downloads
    again = await service.read_file("11", "2", 30_000)
    assert again.pages == "2-2" and moodle.downloads == downloads  # local copy reused

    slides = await service.read_file("12", None, 30_000)
    assert slides.file_type == "powerpoint" and "Agenda" in slides.text


async def test_read_file_next_pages_and_folder_errors(service, moodle):
    part = await service.read_file("11", None, 1000)
    assert not part.truncated and part.next_pages is None

    with pytest.raises(MoodleError) as exc:
        await service.read_file("90", None, 1000)
    assert "90/a.txt" in str(exc.value)
    assert (await service.read_file("90/b.txt", None, 1000)).text == "B"


def test_prompts_are_registered(config):
    server = create_server(config)
    prompts = {p.name for p in asyncio.run(server.list_prompts())}
    assert prompts == {"study_assistant", "last_lecture_summary", "deadlines_overview", "whats_new",
                       "exam_preparation"}
    assert "moodle_read_file" in STUDY_GUIDE
